"""Eco fee report for CRM orders with UPD information."""

from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from html import unescape
from io import BytesIO
import re

from dateutil.relativedelta import relativedelta
from flask import abort, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from sqlalchemy import and_, case, func, literal, or_, select, union_all

from config import settings
from models import (Order, FastOrderCompanies, Shoe, Clothes, Socks, Linen, Parfum,
                    Cosmetics, Toys, ShoeQuantitySize, ClothesQuantitySize,
                    SocksQuantitySize, LinenQuantitySize, LinenSetItem)
from models import db
from tezaurus.runtime_catalogs import get_processing_companies
from views.main.subcategory_notes import COSMETICS_SUBCATEGORY_NOTE_KEYS, TOYS_SUBCATEGORY_NOTE_KEYS
from views.crm.schema import CompaniesOperators


POSITION_RELATIONS = {
    settings.Shoes.CATEGORY: Shoe, settings.Clothes.CATEGORY: Clothes,
    settings.Socks.CATEGORY: Socks, settings.Linen.CATEGORY: Linen,
    settings.Parfum.CATEGORY: Parfum, settings.Cosmetics.CATEGORY: Cosmetics,
    settings.Toys.CATEGORY: Toys,
}


def _quantity_expr(category, model):
    size_models = {
        settings.Shoes.CATEGORY: (ShoeQuantitySize, ShoeQuantitySize.shoe_id),
        settings.Clothes.CATEGORY: (ClothesQuantitySize, ClothesQuantitySize.cl_id),
        settings.Socks.CATEGORY: (SocksQuantitySize, SocksQuantitySize.socks_id),
        settings.Linen.CATEGORY: (LinenQuantitySize, LinenQuantitySize.lin_id),
    }
    if category in size_models:
        size_model, fk = size_models[category]
        total_sizes = (select(func.coalesce(func.sum(size_model.quantity), 0))
                       .where(fk == model.id).correlate(model).scalar_subquery())
        quantity = total_sizes * func.coalesce(model.box_quantity, 1)
        if category == settings.Linen.CATEGORY:
            size_count = (select(func.count(size_model.id)).where(fk == model.id)
                          .correlate(model).scalar_subquery())
            has_set_items = select(literal(1)).where(LinenSetItem.lin_id == model.id).exists()
            return case(
                (has_set_items, func.coalesce(model.kpb_quantity, 1)),
                (model.with_packages == "да", size_count * func.coalesce(model.box_quantity, 1)),
                else_=quantity,
            )
        return quantity
    if category == settings.Parfum.CATEGORY:
        return case((model.with_packages == "да", func.coalesce(model.box_quantity, 0)),
                    else_=func.coalesce(model.quantity, 0))
    return func.coalesce(model.quantity, 0)


def _filters(args):
    today = datetime.now().date()
    def parse_date(value, default):
        if not value:
            return default
        for date_format in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, date_format).date()
            except ValueError:
                continue
        raise ValueError("Неверный формат даты")

    try:
        date_from = parse_date(args.get("date_from"), today.replace(day=1))
        date_to = parse_date(args.get("date_to"), today)
    except ValueError:
        abort(400, "Неверный формат даты")
    if date_from > date_to:
        abort(400, "Начальная дата позже конечной")
    return {
        "date_from": date_from, "date_to": date_to,
        "company": args.get("company", "").strip(),
        "category": args.get("category", "").strip(),
        "subcategory": args.get("subcategory", "").strip(),
        "sort": "asc" if args.get("sort") == "asc" else "desc",
    }


def _positions_query(filters, *, include_company=True):
    """A database query of report positions with all fields needed for display/export."""
    has_structured_upd = and_(Order.processing_company_label.isnot(None),
                              Order.processing_company_label != "",
                              Order.upd_number.isnot(None), Order.upd_number != "")
    company_default = case(
        (has_structured_upd, Order.processing_company_label),
        else_=func.trim(func.split_part(Order.processing_info, "<br>", 1)),
    )
    company_moderation = case(
        (FastOrderCompanies.processing_company_title != "",
         func.concat(FastOrderCompanies.processing_company_title,
                     case((FastOrderCompanies.processing_company_inn != "",
                           func.concat(" (", FastOrderCompanies.processing_company_inn, ")")), else_=""))),
        else_=FastOrderCompanies.processing_company_external_id,
    )
    company = case(
        (Order.is_moderation.is_(True), company_moderation),
        else_=company_default,
    )
    date_start = datetime.combine(filters["date_from"], datetime.min.time())
    date_end = datetime.combine(filters["date_to"] + timedelta(days=1), datetime.min.time())
    base_filters = (
        Order.stage.in_((settings.OrderStage.SENT, settings.OrderStage.CRM_PROCESSED)),
        Order.sent_at >= date_start,
        Order.sent_at < date_end,
        or_(
            and_(Order.is_moderation.is_(False),
                 or_(has_structured_upd, Order.processing_info.contains("УПД:"))),
            and_(Order.is_moderation.is_(True),
                 FastOrderCompanies.upd_number.isnot(None),
                 FastOrderCompanies.upd_number != ""),
        ),
    )
    branches = []
    for category, model in POSITION_RELATIONS.items():
        if filters["category"] and category != filters["category"]:
            continue
        subcategory = getattr(model, "subcategory", literal(""))
        if filters["subcategory"] and not hasattr(model, "subcategory"):
            continue
        stmt = (select(literal(category).label("category"), model.id.label("position_id"),
                       Order.id.label("order_id"), Order.order_idn.label("order_idn"),
                       Order.sent_at.label("sort_date"), Order.is_moderation.label("is_moderation"),
                       company.label("company"),
                       subcategory.label("subcategory"),
                       getattr(model, "article", getattr(model, "model_article", literal(""))).label("article"),
                       model.type.label("product"), model.tnved_code.label("tnved"),
                       _quantity_expr(category, model).label("quantity"),
                       model.weight_grams.label("weight"))
                .select_from(model)
                .join(Order, model.order_id == Order.id)
                .outerjoin(FastOrderCompanies, model.fast_order_company_id == FastOrderCompanies.id)
                .where(Order.category == category, *base_filters))
        if filters["subcategory"]:
            stmt = stmt.where(subcategory == filters["subcategory"])
        branches.append(stmt)
    if not branches:
        return None
    positions = union_all(*branches).subquery("eco_positions")
    if include_company and filters["company"]:
        positions = select(positions).where(positions.c.company == filters["company"]).subquery("eco_filtered_positions")
    return positions


def _position_rows(records):
    """Convert SQL result rows without additional ORM queries."""
    return [{
            "date": record.sort_date.strftime("%d.%m.%Y"),
            "order": record.order_idn or str(record.order_id),
            "company": record.company,
            "category": record.category,
            "subcategory": record.subcategory or "",
            "article": record.article or "",
            "product": record.product or "",
            "tnved": record.tnved or "",
            "quantity": record.quantity or 0,
            "weight": record.weight,
            "is_moderation": record.is_moderation,
        } for record in records]


HEADERS = ("Дата", "Заказ", "Проводящая фирма", "Категория", "Подкатегория",
           "Артикул", "Вид товара", "ТН ВЭД", "Количество марок", "Вес позиции, г")
KEYS = ("date", "order", "company", "category", "subcategory", "article", "product", "tnved", "quantity", "weight")


def _missing_weight_expr(positions):
    """Only missing weight in ordinary orders makes a report total incomplete."""
    return case((and_(positions.c.weight.is_(None), positions.c.is_moderation.is_(False)), 1), else_=0)


COMPANY_CATEGORY_LABELS = {
    settings.Clothes.CATEGORY: "Одежда",
    settings.Shoes.CATEGORY: "Обувь",
    settings.Parfum.CATEGORY: "Духи",
    settings.Cosmetics.CATEGORY: "Косметика",
}
COMPANY_INN_RE = re.compile(r"(?<!\d)\d{10}(?:\d{2})?(?!\d)")


def _company_identity(label):
    label = (label or "").strip()
    match = COMPANY_INN_RE.search(label)
    if match:
        return match.group()
    return " ".join(label.replace('"', "").replace("«", "").replace("»", "").lower().split())


def _company_summary(positions, filters):
    """Summarize all catalog companies, including those with no report positions."""
    catalog = get_processing_companies()
    catalog.extend({"title": item.display_name, "inn": item.inn} for item in CompaniesOperators)
    companies = {}
    aliases = {}
    for item in catalog:
        name = (item.get("title") or "").strip()
        inn = (item.get("inn") or "").strip()
        key = inn or _company_identity(name)
        if key and key not in companies:
            companies[key] = {"name": name or inn, "positions": 0, "quantity": 0, "weight": Decimal(0),
                              "known_weights": 0, "missing": 0, "categories": {}}
        if key:
            aliases[_company_identity(name)] = key
            if item.get("external_id"):
                aliases[str(item["external_id"])] = key

    if positions is not None:
        summary_query = (select(positions.c.company, positions.c.category,
                                func.count().label("positions"),
                                func.coalesce(func.sum(positions.c.quantity), 0).label("quantity"),
                                func.coalesce(func.sum(positions.c.weight), 0).label("weight"),
                                func.count(positions.c.weight).label("known_weights"),
                                func.sum(_missing_weight_expr(positions)).label("missing"))
                         .group_by(positions.c.company, positions.c.category))
        for row in db.session.execute(summary_query):
            identity = _company_identity(row.company)
            key = aliases.get(identity, identity) or "Без фирмы"
            if key not in companies:
                companies[key] = {"name": row.company or "Без фирмы", "positions": 0, "quantity": 0,
                                  "weight": Decimal(0), "known_weights": 0, "missing": 0, "categories": {}}
            company = companies[key]
            category = company["categories"].setdefault(row.category, {
                "name": COMPANY_CATEGORY_LABELS.get(row.category, row.category.capitalize()),
                "positions": 0, "quantity": 0, "weight": Decimal(0), "known_weights": 0, "missing": 0,
            })
            positions_count = int(row.positions or 0)
            quantity = int(row.quantity or 0)
            weight = Decimal(row.weight or 0)
            known_weights = int(row.known_weights or 0)
            missing = int(row.missing or 0)
            company["positions"] += positions_count
            company["quantity"] += quantity
            company["weight"] += weight
            company["known_weights"] += known_weights
            company["missing"] += missing
            category["positions"] += positions_count
            category["quantity"] += quantity
            category["weight"] += weight
            category["known_weights"] += known_weights
            category["missing"] += missing

    if filters["company"]:
        identity = _company_identity(filters["company"])
        selected = aliases.get(identity, identity)
        companies = {key: value for key, value in companies.items() if key == selected}
        if not companies:
            companies[selected] = {"name": filters["company"], "positions": 0, "quantity": 0,
                                   "weight": Decimal(0), "known_weights": 0, "missing": 0, "categories": {}}

    def weight_values(item):
        if item["missing"] or not item["known_weights"]:
            return None, "—"
        tonnes = item["weight"] / Decimal(1000000)
        label = f"{tonnes:,.6f}".rstrip("0").rstrip(".").replace(",", " ").replace(".", ",")
        return float(tonnes), label

    result = []
    for company in sorted(companies.values(), key=lambda item: (-item["quantity"], item["name"].casefold())):
        categories = sorted(company["categories"].values(), key=lambda item: (-item["quantity"], item["name"].casefold()))
        weight_tonnes, weight_label = weight_values(company)
        category_rows = []
        for category in categories:
            category_weight, category_label = weight_values(category)
            category_rows.append({"name": category["name"], "positions": category["positions"],
                                  "quantity": category["quantity"],
                                  "weight_tonnes": category_weight, "weight_label": category_label,
                                  "missing": category["missing"]})
        result.append({"name": company["name"], "positions": company["positions"],
                       "quantity": company["quantity"],
                       "weight_tonnes": weight_tonnes, "weight_label": weight_label,
                       "missing": company["missing"],
                       "categories": category_rows})
    return result


def _subcategory_label(value):
    if value == "common":
        return "основная одежда"
    return settings.SUB_CATEGORIES_DICT.get(value, value)


def eco_sbor_report():
    filters = _filters(request.args)
    report_format = request.args.get("format")
    if report_format not in ("xlsx", "json"):
        subcategories = settings.SUB_CATEGORIES_DICT
        subcategory_categories = {
            key: (settings.Cosmetics.CATEGORY if key in COSMETICS_SUBCATEGORY_NOTE_KEYS else
                  settings.Toys.CATEGORY if key in TOYS_SUBCATEGORY_NOTE_KEYS else settings.Clothes.CATEGORY)
            for key in subcategories
        }
        return render_template("admin/eco_sbor_report.html", filters=filters,
                               categories=POSITION_RELATIONS, subcategories=subcategories,
                               subcategory_categories=subcategory_categories,
                               company_label=unescape(filters["company"]))

    if report_format == "xlsx" and filters["date_to"] > filters["date_from"] + relativedelta(months=12):
        return jsonify(message="Для Excel выберите период не более 12 месяцев."), 400

    positions = _positions_query(filters)
    company_summary = _company_summary(positions, filters)
    order_direction = (positions.c.sort_date.asc() if filters["sort"] == "asc" else positions.c.sort_date.desc()) if positions is not None else None
    if report_format == "xlsx":
        book = Workbook(write_only=True)
        company_sheet = book.create_sheet("Фирмы")
        company_sheet.freeze_panes = "A2"
        for column, width in {"A": 8, "B": 48, "C": 23, "D": 20, "E": 24}.items():
            company_sheet.column_dimensions[column].width = width
        headers = []
        for title in ("№", "Фирма / категория", "Количество КМ", "Вес, т", "Вес не указан у позиций"):
            cell = WriteOnlyCell(company_sheet, value=title)
            cell.font = Font(bold=True)
            headers.append(cell)
        company_sheet.append(headers)
        for index, company in enumerate(company_summary, 1):
            company_cells = []
            for column, value in enumerate((index, company["name"], company["quantity"],
                                            company["weight_tonnes"] if company["weight_tonnes"] is not None else "—",
                                            company["missing"]), 1):
                cell = WriteOnlyCell(company_sheet, value=value)
                cell.font = Font(bold=True)
                if column == 3:
                    cell.number_format = '#,##0'
                elif column == 4 and isinstance(value, (float, int)):
                    cell.number_format = '#,##0.######'
                company_cells.append(cell)
            company_sheet.append(company_cells)
            for category in company["categories"]:
                category_cells = []
                for column, value in enumerate((None, f"    {category['name']}", category["quantity"],
                                                category["weight_tonnes"] if category["weight_tonnes"] is not None else "—",
                                                category["missing"]), 1):
                    cell = WriteOnlyCell(company_sheet, value=value)
                    if column == 3:
                        cell.number_format = '#,##0'
                    elif column == 4 and isinstance(value, (float, int)):
                        cell.number_format = '#,##0.######'
                    category_cells.append(cell)
                company_sheet.append(category_cells)
        sheets = {}
        sheet_lengths = {}
        sheet_totals = defaultdict(lambda: {"quantity": 0, "weight": 0, "known_weights": 0, "missing": 0})
        if positions is not None:
            stmt = select(positions).order_by(positions.c.category, positions.c.subcategory,
                                               order_direction, positions.c.position_id)
            result = db.session.execute(stmt.execution_options(stream_results=True))
            while records := result.fetchmany(500):
                for item in _position_rows(records):
                    key = (item["category"], item["subcategory"])
                    if key not in sheets:
                        title = f"{len(sheets) + 1} {key[0]} {_subcategory_label(key[1])}"[:31]
                        sheet = book.create_sheet(title)
                        sheet.freeze_panes = "A2"
                        for column, width in {"A": 17, "B": 18, "C": 42, "D": 20, "E": 28, "F": 25, "G": 34, "H": 18, "I": 15, "J": 20}.items():
                            sheet.column_dimensions[column].width = width
                        headers = []
                        for header in HEADERS:
                            cell = WriteOnlyCell(sheet, value=header)
                            cell.font = Font(bold=True)
                            headers.append(cell)
                        sheet.append(headers)
                        sheets[key] = sheet
                        sheet_lengths[key] = 1
                    sheets[key].append([_subcategory_label(item[column]) if column == "subcategory" else item[column]
                                        for column in KEYS])
                    sheet_lengths[key] += 1
                    sheet_totals[key]["quantity"] += item["quantity"]
                    if item["weight"] is None:
                        if not item["is_moderation"]:
                            sheet_totals[key]["missing"] += 1
                    else:
                        sheet_totals[key]["weight"] += item["weight"]
                        sheet_totals[key]["known_weights"] += 1
        for key, sheet in sheets.items():
            sheet.auto_filter.ref = f"A1:J{sheet_lengths[key]}"
            totals = sheet_totals[key]
            footer = ["ИТОГО", None, None, None, None, None, None, None,
                      totals["quantity"], "—" if totals["missing"] or not totals["known_weights"] else totals["weight"]]
            footer_cells = []
            for column, value in enumerate(footer, 1):
                cell = WriteOnlyCell(sheet, value=value)
                cell.font = Font(bold=True)
                if column == 9:
                    cell.number_format = '#,##0" шт."'
                elif column == 10:
                    cell.number_format = '#,##0" г"'
                footer_cells.append(cell)
            sheet.append(footer_cells)
        output = BytesIO()
        book.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="eco_sbor.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    groups = {}
    companies = [] if request.args.get("companies") != "0" else None
    if positions is not None:
        group_query = (select(positions.c.category, positions.c.subcategory,
                              func.count().label("position_count"),
                              func.coalesce(func.sum(positions.c.quantity), 0).label("quantity_total"),
                              func.coalesce(func.sum(positions.c.weight), 0).label("weight_total"),
                              func.count(positions.c.weight).label("known_weights"),
                              func.sum(_missing_weight_expr(positions)).label("missing_weight"))
                       .group_by(positions.c.category, positions.c.subcategory)
                       .order_by(positions.c.category, positions.c.subcategory))
        groups = {(row.category, row.subcategory or ""):
                  {"count": row.position_count, "quantity": row.quantity_total,
                   "weight": row.weight_total, "known_weights": row.known_weights,
                   "missing_weight": row.missing_weight}
                  for row in db.session.execute(group_query)}
        if request.args.get("companies") != "0":
            unfiltered_positions = _positions_query(filters, include_company=False)
            company_names = db.session.execute(
                select(unfiltered_positions.c.company).distinct().order_by(unfiltered_positions.c.company)
            ).scalars().all()
            companies = [{"value": name, "label": unescape(name or "")}
                         for name in company_names]
    keys = list(groups)
    try:
        group_index = int(request.args.get("group", -1))
    except ValueError:
        group_index = -1
    if group_index < -1 or group_index >= len(keys):
        group_index = -1
    group_stats = groups[keys[group_index]] if group_index >= 0 else {"count": 0, "quantity": 0, "weight": None, "known_weights": 0, "missing_weight": 0}
    group_total = group_stats["count"]
    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1
    per_page = settings.PAGINATION_PER_PAGE
    page = min(page, max(1, (group_total + per_page - 1) // per_page))
    offset = (page - 1) * per_page
    items = []
    if group_index >= 0:
        category, subcategory = keys[group_index]
        stmt = (select(positions)
                .where(positions.c.category == category, positions.c.subcategory == subcategory)
                .order_by(order_direction, positions.c.position_id)
                .offset(offset).limit(per_page))
        items = _position_rows(db.session.execute(stmt).all())
    html = render_template("admin/eco_sbor_table.html", filters=filters, groups=groups,
                           keys=keys, group_index=group_index, items=items,
                           company_summary=company_summary,
                           group_stats=group_stats, group_total=group_total, offset=offset, page=page,
                           per_page=per_page,
                           total=sum(group["count"] for group in groups.values()),
                           total_quantity=sum(group["quantity"] for group in groups.values()),
                           total_weight=sum(group["weight"] or 0 for group in groups.values()) if any(group["weight"] is not None for group in groups.values()) else None,
                           total_known_weights=sum(group["known_weights"] for group in groups.values()),
                           total_missing=sum(group["missing_weight"] for group in groups.values()),
                           subcategory_names=settings.SUB_CATEGORIES_DICT)
    return jsonify(htmlresponse=html, companies=companies)
